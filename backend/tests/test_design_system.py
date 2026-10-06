import io
import re
import docx
import pypdf
import pytest
from app.services.export_service import generate_resume_pdf, generate_resume_docx

test_resume = {
    "candidate_name": "Sarah Connor",
    "headline": "VP of Engineering",
    "contact_info": {
        "email": "sarah@example.com",
        "phone": "+1 555 0199",
        "location": "Los Angeles, CA",
        "linkedin": "linkedin.com/in/sarahconnor",
        "github": "github.com/sarahconnor"
    },
    "summary": 'Executive leader specializing in <span style="color: #e11d48;">mission-critical AI defense systems</span> and team scaling.',
    "experience": [
        {
            "role": "VP of Engineering",
            "company": "Cyberdyne Systems",
            "location": "Los Angeles, CA",
            "start_date": "2019",
            "end_date": "Present",
            "bullets": [
                'Scaled engineering organization from <font color="#2563eb">15 to 140</font> distributed engineers.',
                "Delivered 99.999% reliability SLA across all deployments."
            ]
        }
    ],
    "skills": [
        {"category": "Architecture", "skills": "Distributed Systems, Microservices, Cloud Security"},
        {"category": "Leadership", "skills": "Strategic Planning, Team Scaling, Budgeting"}
    ],
    "education": [
        {"degree": "B.S. in Computer Science", "institution": "Stanford University", "start_date": "2010", "end_date": "2014"}
    ],
    "certifications": [
        {"name": "AWS Certified Solutions Architect - Professional"}
    ],
    "languages": [
        {"language": "English", "proficiency": "Native"}
    ]
}

dummy_photo = "data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="

TEMPLATES = [
    "classic_ats",
    "modern_professional",
    "minimal_professional",
    "technical_ats",
    "executive_professional",
    "two_column_professional",
    "creative_professional"
]

@pytest.mark.parametrize("tpl", TEMPLATES)
def test_all_templates_pdf_and_docx(tpl):
    pdf = generate_resume_pdf(test_resume, template_name=tpl, font_family="Inter", page_size="a4", margins="standard")
    assert len(pdf) > 1000
    reader = pypdf.PdfReader(io.BytesIO(pdf))
    assert len(reader.pages) >= 1

    doc_bytes = generate_resume_docx(test_resume, template_name=tpl, font_family="Inter", page_size="a4", margins="standard")
    assert len(doc_bytes) > 2000
    doc = docx.Document(io.BytesIO(doc_bytes))
    assert len(doc.paragraphs) > 5

def test_two_column_ats_reading_order():
    pdf = generate_resume_pdf(test_resume, template_name="two_column_professional", layout="two_column")
    reader = pypdf.PdfReader(io.BytesIO(pdf))
    raw_text = " ".join(page.extract_text() for page in reader.pages)
    norm_text = re.sub(r"\s+", " ", raw_text)

    idx_name = norm_text.find("Sarah Connor")
    idx_exp = norm_text.find("Cyberdyne Systems")
    idx_cert = norm_text.find("AWS Certified Solutions Architect")

    assert idx_name != -1, "Name missing"
    assert idx_exp != -1, "Experience missing"
    assert idx_cert != -1, "Certifications missing"
    assert idx_name < idx_exp, "Name must precede experience"
    assert idx_exp < idx_cert, "Primary column must precede secondary column in reading order"

def test_technical_ats_clean_design():
    pdf = generate_resume_pdf(test_resume, template_name="technical_ats")
    reader = pypdf.PdfReader(io.BytesIO(pdf))
    text = " ".join(page.extract_text() for page in reader.pages)
    assert "//" not in text, "Found legacy // heading artifact in technical_ats"
    assert "terminal" not in text.lower(), "Found terminal references in technical_ats"

def test_pdf_page_size_and_customizations():
    pdf = generate_resume_pdf(
        test_resume,
        template_name="executive_professional",
        accent_color="#7c3aed",
        secondary_color="#475569",
        font_family="Georgia",
        page_size="letter",
        margins="narrow",
        line_height="comfortable",
        photo_url=dummy_photo,
        photo_shape="rounded",
        photo_size="md",
        element_colors={
            "name": "#581c87",
            "headline": "#7e22ce",
            "body": "#1e293b",
            "dividers": "#c084fc",
            "bullets": "#9333ea"
        },
        section_styles={
            "experiences": {
                "titleColor": "#6b21a8",
                "dividerStyle": "double",
                "dividerColor": "#a855f7",
                "alignment": "left"
            }
        }
    )
    assert len(pdf) > 1000
    reader = pypdf.PdfReader(io.BytesIO(pdf))
    p0 = reader.pages[0]
    # Letter size is 612 x 792 pt
    assert p0.mediabox.width == 612.0
    assert p0.mediabox.height == 792.0

def test_docx_page_size_and_margins():
    docx_bytes = generate_resume_docx(
        test_resume,
        template_name="executive_professional",
        accent_color="#7c3aed",
        font_family="Georgia",
        page_size="letter",
        margins="narrow",
        photo_url=dummy_photo
    )
    assert len(docx_bytes) > 2000
    doc = docx.Document(io.BytesIO(docx_bytes))
    sec = doc.sections[0]
    assert abs(sec.page_width.inches - 8.5) < 0.01
    assert abs(sec.page_height.inches - 11.0) < 0.01
    assert abs(sec.left_margin.inches - 0.35) < 0.01
