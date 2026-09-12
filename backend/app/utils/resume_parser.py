from io import BytesIO
from pathlib import Path

from docx import Document
from fastapi import UploadFile, status
from pypdf import PdfReader

from app.core.config import get_settings
from app.core.errors import AppError
from app.utils.sanitize import clean_text


ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}


async def extract_upload_text(upload: UploadFile) -> str:
    settings = get_settings()
    content = await upload.read()
    if len(content) > settings.upload_max_bytes:
        raise AppError("Uploaded file is too large.", status.HTTP_413_REQUEST_ENTITY_TOO_LARGE)

    extension = Path(upload.filename or "").suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise AppError("Upload must be a PDF, DOCX, or TXT file.", status.HTTP_400_BAD_REQUEST)

    if extension == ".pdf":
        return clean_text(_extract_pdf_text(content))
    if extension == ".docx":
        return clean_text(_extract_docx_text(content))
    return clean_text(content.decode("utf-8", errors="ignore"))


def _extract_pdf_text(content: bytes) -> str:
    reader = PdfReader(BytesIO(content))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages)


def _extract_docx_text(content: bytes) -> str:
    document = Document(BytesIO(content))
    parts: list[str] = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)
