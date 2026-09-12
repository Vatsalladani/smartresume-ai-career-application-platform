import difflib
import re
from io import BytesIO
from typing import Any

from docx import Document
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
from sqlalchemy.orm import Session

from app.models import Resume, ResumeVersion
from app.utils.sanitize import clean_text


SECTION_HINTS = {
    "summary": ["summary", "profile", "objective"],
    "skills": ["skills", "technical skills", "core skills"],
    "experience": ["experience", "work experience", "employment"],
    "education": ["education", "academic"],
    "projects": ["projects", "portfolio"],
}


def parse_resume_text(text: str) -> dict[str, Any]:
    cleaned = clean_text(text)
    lines = [line.strip() for line in cleaned.splitlines() if line.strip()]
    sections: dict[str, list[str]] = {key: [] for key in SECTION_HINTS}
    current = "summary"

    for line in lines:
        lower = line.lower().strip(":")
        matched = next((key for key, labels in SECTION_HINTS.items() if lower in labels), None)
        if matched:
            current = matched
            continue
        sections.setdefault(current, []).append(line)

    email_match = re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", cleaned)
    phone_match = re.search(r"(?:\+?\d[\d\s().-]{7,}\d)", cleaned)
    return {
        "contact": {
            "email": email_match.group(0) if email_match else None,
            "phone": phone_match.group(0) if phone_match else None,
        },
        "sections": sections,
        "line_count": len(lines),
        "word_count": len(re.findall(r"\w+", cleaned)),
    }


def calculate_completeness(parsed: dict[str, Any], text: str) -> int:
    score = 0
    contact = parsed.get("contact", {})
    sections = parsed.get("sections", {})
    if contact.get("email"):
        score += 15
    if contact.get("phone"):
        score += 10
    if sections.get("summary"):
        score += 15
    if sections.get("skills"):
        score += 20
    if sections.get("experience"):
        score += 25
    if sections.get("education"):
        score += 10
    if re.search(r"\d+%|\$\d+|\b\d+x\b|\b\d+\+", text.lower()):
        score += 5
    return min(score, 100)


def create_resume(db: Session, *, user_id: int, title: str, raw_text: str, changelog: str = "Initial import") -> Resume:
    parsed = parse_resume_text(raw_text)
    resume = Resume(
        user_id=user_id,
        title=title.strip() or "My Resume",
        raw_text=clean_text(raw_text),
        parsed_content=parsed,
        completeness_score=calculate_completeness(parsed, raw_text),
    )
    db.add(resume)
    db.flush()
    add_version(db, resume, changelog=changelog)
    return resume


def add_version(db: Session, resume: Resume, changelog: str, analysis_snapshot: dict | None = None) -> ResumeVersion:
    latest = (
        db.query(ResumeVersion)
        .filter(ResumeVersion.resume_id == resume.id)
        .order_by(ResumeVersion.version_number.desc())
        .first()
    )
    version = ResumeVersion(
        resume_id=resume.id,
        version_number=(latest.version_number + 1 if latest else 1),
        source_text=resume.raw_text,
        content=resume.parsed_content,
        analysis_snapshot=analysis_snapshot,
        changelog=changelog,
    )
    db.add(version)
    return version


def update_resume(db: Session, resume: Resume, *, title: str | None = None, raw_text: str | None = None) -> Resume:
    if title is not None:
        resume.title = title.strip()
    if raw_text is not None:
        resume.raw_text = clean_text(raw_text)
        resume.parsed_content = parse_resume_text(raw_text)
        resume.completeness_score = calculate_completeness(resume.parsed_content, raw_text)
        add_version(db, resume, changelog="Manual edit")
    return resume


def duplicate_resume(db: Session, resume: Resume, user_id: int) -> Resume:
    return create_resume(
        db,
        user_id=user_id,
        title=f"{resume.title} Copy",
        raw_text=resume.raw_text,
        changelog=f"Duplicated from resume {resume.id}",
    )


def restore_version(db: Session, resume: Resume, version: ResumeVersion) -> Resume:
    resume.raw_text = version.source_text
    resume.parsed_content = version.content
    resume.completeness_score = calculate_completeness(resume.parsed_content, resume.raw_text)
    add_version(db, resume, changelog=f"Restored version {version.version_number}")
    return resume


def compare_with_version(resume: Resume, version: ResumeVersion) -> dict[str, Any]:
    current_lines = resume.raw_text.splitlines()
    old_lines = version.source_text.splitlines()
    diff = list(difflib.ndiff(old_lines, current_lines))
    return {
        "resume_id": resume.id,
        "current_version": len(resume.versions),
        "compared_version": version.version_number,
        "additions": [line[2:] for line in diff if line.startswith("+ ")],
        "removals": [line[2:] for line in diff if line.startswith("- ")],
    }


def export_resume_pdf(resume: Resume, analysis: dict | None = None) -> bytes:
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=LETTER, rightMargin=54, leftMargin=54, topMargin=54, bottomMargin=54)
    styles = getSampleStyleSheet()
    story = [Paragraph(resume.title, styles["Title"]), Spacer(1, 12)]
    for paragraph in resume.raw_text.splitlines():
        if paragraph.strip():
            story.append(Paragraph(paragraph.strip(), styles["BodyText"]))
            story.append(Spacer(1, 6))
    if analysis:
        story.append(Spacer(1, 12))
        story.append(Paragraph("ATS Summary", styles["Heading2"]))
        story.append(Paragraph(f"Predicted score: {analysis.get('predicted_ats_score', 0)}/100", styles["BodyText"]))
    doc.build(story)
    return buffer.getvalue()


def export_resume_docx(resume: Resume, analysis: dict | None = None) -> bytes:
    document = Document()
    document.add_heading(resume.title, level=1)
    for paragraph in resume.raw_text.splitlines():
        if paragraph.strip():
            document.add_paragraph(paragraph.strip())
    if analysis:
        document.add_heading("ATS Summary", level=2)
        document.add_paragraph(f"Predicted score: {analysis.get('predicted_ats_score', 0)}/100")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
