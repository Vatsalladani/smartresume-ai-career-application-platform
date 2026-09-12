from io import BytesIO
from typing import Any
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, HRFlowable
from reportlab.lib import colors

TEMPLATE_COLORS: dict[str, str] = {
    "classic_ats": "#1e3a8a",
    "campus_fresher": "#0d9488",
    "clean_professional": "#334155",
    "technical_ats": "#0284c7",
    "business_professional": "#1e40af",
    "finance_professional": "#15803d",
    "healthcare_pharmacy": "#059669",
    "consulting_management": "#4338ca",
    "experienced_professional": "#374151",
    "academic_research": "#7c2d12",
    "creative_professional": "#7c3aed",
    "executive": "#0f172a",
    "professional": "#1e3a8a",  # Backward compatibility alias
}

TEMPLATE_SECTION_ORDERS: dict[str, list[str]] = {
    "classic_ats": ["summary", "experiences", "education", "skills", "projects", "certifications"],
    "campus_fresher": ["summary", "education", "projects", "skills", "experiences", "certifications"],
    "clean_professional": ["summary", "experiences", "skills", "education", "projects", "certifications"],
    "technical_ats": ["skills", "projects", "experiences", "education", "certifications", "summary"],
    "business_professional": ["summary", "experiences", "projects", "education", "skills", "certifications"],
    "finance_professional": ["summary", "experiences", "skills", "education", "certifications", "projects"],
    "healthcare_pharmacy": ["summary", "certifications", "experiences", "education", "skills"],
    "consulting_management": ["summary", "experiences", "projects", "skills", "education", "certifications"],
    "experienced_professional": ["summary", "experiences", "skills", "projects", "education", "certifications"],
    "academic_research": ["summary", "education", "projects", "experiences", "skills", "certifications"],
    "creative_professional": ["summary", "projects", "experiences", "skills", "education", "certifications"],
    "executive": ["summary", "experiences", "projects", "education", "certifications", "skills"],
    "professional": ["summary", "experiences", "skills", "education", "projects", "certifications"],
}


def build_pdf_styles(accent_hex: str = "#1e3a8a"):
    accent = colors.HexColor(accent_hex)
    text_color = colors.HexColor("#1f2937")
    muted_color = colors.HexColor("#4b5563")

    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        name="CandidateName",
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        textColor=accent,
        alignment=0,
    ))
    styles.add(ParagraphStyle(
        name="CandidateHeadline",
        fontName="Helvetica",
        fontSize=11,
        leading=14,
        textColor=muted_color,
        alignment=0,
    ))
    styles.add(ParagraphStyle(
        name="ContactBar",
        fontName="Helvetica",
        fontSize=9,
        leading=12,
        textColor=muted_color,
        alignment=0,
    ))
    styles.add(ParagraphStyle(
        name="SectionHeader",
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=16,
        textColor=accent,
        spaceBefore=8,
        spaceAfter=4,
    ))
    styles.add(ParagraphStyle(
        name="ItemTitle",
        fontName="Helvetica-Bold",
        fontSize=10,
        leading=13,
        textColor=text_color,
    ))
    styles.add(ParagraphStyle(
        name="ItemSub",
        fontName="Helvetica-Oblique",
        fontSize=9,
        leading=12,
        textColor=muted_color,
    ))
    styles.add(ParagraphStyle(
        name="ResumeBullet",
        fontName="Helvetica",
        fontSize=9.5,
        leading=13,
        textColor=text_color,
        leftIndent=14,
        firstLineIndent=-10,
        spaceAfter=2,
    ))
    styles.add(ParagraphStyle(
        name="ResumeBody",
        fontName="Helvetica",
        fontSize=9.5,
        leading=13,
        textColor=text_color,
        spaceAfter=4,
    ))
    return styles


def generate_resume_pdf(content: dict[str, Any], template_name: str = "classic_ats") -> bytes:
    buffer = BytesIO()
    # 0.5 inch margins for standard ATS parsing
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36,
    )
    accent_hex = TEMPLATE_COLORS.get(template_name, "#1e3a8a")
    styles = build_pdf_styles(accent_hex)
    story = []

    name = content.get("candidate_name", "Candidate")
    headline = content.get("headline", "")
    summary = content.get("summary", "")
    phone = content.get("phone", "")
    location = content.get("location", "")
    linkedin = content.get("linkedin_url", "")
    github = content.get("github_url", "")
    website = content.get("website_url", "")

    contact_parts = [p for p in [phone, location, linkedin, github, website] if p]
    contact_str = " | ".join(contact_parts)

    # Header section
    story.append(Paragraph(name, styles["CandidateName"]))
    if headline:
        story.append(Paragraph(headline, styles["CandidateHeadline"]))
    if contact_str:
        story.append(Paragraph(contact_str, styles["ContactBar"]))
    story.append(Spacer(1, 4))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor(accent_hex), spaceBefore=4, spaceAfter=8))

    experiences = content.get("experiences", [])
    projects = content.get("projects", [])
    skills = content.get("skills", [])
    education = content.get("education", [])
    certifications = content.get("certifications", [])

    def render_summary():
        if summary:
            title = "EXECUTIVE SUMMARY" if template_name == "executive" else "PROFESSIONAL SUMMARY"
            story.append(Paragraph(title, styles["SectionHeader"]))
            story.append(Paragraph(summary, styles["ResumeBody"]))
            story.append(Spacer(1, 4))

    def render_skills():
        if skills:
            if template_name == "technical_ats":
                header_title = "CORE TECHNICAL SKILLS & ARCHITECTURE"
            elif template_name == "finance_professional":
                header_title = "FINANCIAL & QUANTITATIVE COMPETENCIES"
            elif template_name == "healthcare_pharmacy":
                header_title = "CLINICAL COMPETENCIES"
            elif template_name == "executive":
                header_title = "EXECUTIVE LEADERSHIP & CORE COMPETENCIES"
            else:
                header_title = "SKILLS"
            story.append(Paragraph(header_title, styles["SectionHeader"]))
            skill_text = " • ".join(skills if isinstance(skills[0], str) else [s.get("name", "") for s in skills])
            story.append(Paragraph(skill_text, styles["ResumeBody"]))
            story.append(Spacer(1, 4))

    def render_experiences():
        if experiences:
            header_title = "ENGAGEMENT & ADVISORY EXPERIENCE" if template_name == "consulting_management" else "WORK EXPERIENCE"
            story.append(Paragraph(header_title, styles["SectionHeader"]))
            for exp in experiences:
                comp = exp.get("company", "")
                role = exp.get("role_title", "")
                start = exp.get("start_date", "")
                end = exp.get("end_date", "Present")
                date_str = f"{start} – {end}" if start else ""
                story.append(Paragraph(f"<b>{role}</b> — {comp}", styles["ItemTitle"]))
                if date_str:
                    story.append(Paragraph(date_str, styles["ItemSub"]))
                bullets = exp.get("bullet_points", [])
                for b in bullets:
                    story.append(Paragraph(f"• {b}", styles["ResumeBullet"]))
                story.append(Spacer(1, 3))

    def render_projects():
        if projects:
            header_title = "PUBLICATIONS & RESEARCH" if template_name == "academic_research" else "KEY PROJECTS"
            story.append(Paragraph(header_title, styles["SectionHeader"]))
            for proj in projects:
                title = proj.get("title", "")
                tech = proj.get("technologies", [])
                tech_str = f" ({', '.join(tech)})" if tech else ""
                story.append(Paragraph(f"<b>{title}</b>{tech_str}", styles["ItemTitle"]))
                bullets = proj.get("bullet_points", [])
                for b in bullets:
                    story.append(Paragraph(f"• {b}", styles["ResumeBullet"]))
                story.append(Spacer(1, 3))

    def render_education():
        if education:
            story.append(Paragraph("EDUCATION", styles["SectionHeader"]))
            for edu in education:
                inst = edu.get("institution", "")
                deg = edu.get("degree", "")
                field = edu.get("field_of_study", "")
                start = edu.get("start_date", "")
                end = edu.get("end_date", "")
                date_str = f"{start} – {end}" if start or end else ""
                full_deg = f"{deg} in {field}" if field else deg
                story.append(Paragraph(f"<b>{full_deg}</b> — {inst}", styles["ItemTitle"]))
                if date_str:
                    story.append(Paragraph(date_str, styles["ItemSub"]))
                story.append(Spacer(1, 2))

    def render_certifications():
        if certifications:
            header_title = "LICENSURES & CERTIFICATIONS" if template_name == "healthcare_pharmacy" else "CERTIFICATIONS"
            story.append(Paragraph(header_title, styles["SectionHeader"]))
            for cert in certifications:
                cname = cert.get("name", "") if isinstance(cert, dict) else str(cert)
                issuer = cert.get("issuer", "") if isinstance(cert, dict) else ""
                date = cert.get("issue_date", "") if isinstance(cert, dict) else ""
                meta = f" — {issuer}" if issuer else ""
                if date:
                    meta += f" ({date})"
                story.append(Paragraph(f"<b>{cname}</b>{meta}", styles["ItemTitle"]))
            story.append(Spacer(1, 3))

    section_renderers = {
        "summary": render_summary,
        "skills": render_skills,
        "experiences": render_experiences,
        "projects": render_projects,
        "education": render_education,
        "certifications": render_certifications,
    }

    order = TEMPLATE_SECTION_ORDERS.get(template_name, TEMPLATE_SECTION_ORDERS["classic_ats"])
    for sec in order:
        renderer = section_renderers.get(sec)
        if renderer:
            renderer()

    doc.build(story)
    return buffer.getvalue()


def generate_resume_docx(content: dict[str, Any], template_name: str = "classic_ats") -> bytes:
    doc = Document()
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)

    name = content.get("candidate_name", "Candidate")
    headline = content.get("headline", "")
    summary = content.get("summary", "")
    phone = content.get("phone", "")
    location = content.get("location", "")
    linkedin = content.get("linkedin_url", "")
    github = content.get("github_url", "")
    website = content.get("website_url", "")

    # Title / Header
    h1 = doc.add_heading(name, level=1)
    h1.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if headline:
        p_head = doc.add_paragraph(headline)
        p_head.style.font.italic = True

    contact_parts = [p for p in [phone, location, linkedin, github, website] if p]
    if contact_parts:
        doc.add_paragraph(" | ".join(contact_parts))

    experiences = content.get("experiences", [])
    projects = content.get("projects", [])
    skills = content.get("skills", [])
    education = content.get("education", [])
    certifications = content.get("certifications", [])

    def docx_summary():
        if summary:
            title = "Executive Summary" if template_name == "executive" else "Professional Summary"
            doc.add_heading(title, level=2)
            doc.add_paragraph(summary)

    def docx_skills():
        if skills:
            if template_name == "technical_ats":
                header_title = "Core Technical Skills & Architecture"
            elif template_name == "finance_professional":
                header_title = "Financial & Quantitative Competencies"
            elif template_name == "healthcare_pharmacy":
                header_title = "Clinical Competencies"
            elif template_name == "executive":
                header_title = "Executive Leadership & Core Competencies"
            else:
                header_title = "Skills"
            doc.add_heading(header_title, level=2)
            skill_text = " • ".join(skills if isinstance(skills[0], str) else [s.get("name", "") for s in skills])
            doc.add_paragraph(skill_text)

    def docx_experiences():
        if experiences:
            header_title = "Engagement & Advisory Experience" if template_name == "consulting_management" else "Work Experience"
            doc.add_heading(header_title, level=2)
            for exp in experiences:
                comp = exp.get("company", "")
                role = exp.get("role_title", "")
                start = exp.get("start_date", "")
                end = exp.get("end_date", "Present")
                date_str = f" ({start} – {end})" if start else ""
                p = doc.add_paragraph()
                r1 = p.add_run(f"{role} — {comp}")
                r1.bold = True
                if date_str:
                    r2 = p.add_run(date_str)
                    r2.italic = True
                for b in exp.get("bullet_points", []):
                    doc.add_paragraph(b, style="List Bullet")

    def docx_projects():
        if projects:
            header_title = "Publications & Research" if template_name == "academic_research" else "Key Projects"
            doc.add_heading(header_title, level=2)
            for proj in projects:
                title = proj.get("title", "")
                tech = proj.get("technologies", [])
                p = doc.add_paragraph()
                r1 = p.add_run(title)
                r1.bold = True
                if tech:
                    p.add_run(f" [{', '.join(tech)}]")
                for b in proj.get("bullet_points", []):
                    doc.add_paragraph(b, style="List Bullet")

    def docx_education():
        if education:
            doc.add_heading("Education", level=2)
            for edu in education:
                inst = edu.get("institution", "")
                deg = edu.get("degree", "")
                field = edu.get("field_of_study", "")
                start = edu.get("start_date", "")
                end = edu.get("end_date", "")
                p = doc.add_paragraph()
                full_deg = f"{deg} in {field}" if field else deg
                r1 = p.add_run(f"{full_deg} — {inst}")
                r1.bold = True
                if start or end:
                    p.add_run(f" ({start} – {end})")

    def docx_certifications():
        if certifications:
            header_title = "Licensures & Certifications" if template_name == "healthcare_pharmacy" else "Certifications"
            doc.add_heading(header_title, level=2)
            for cert in certifications:
                cname = cert.get("name", "") if isinstance(cert, dict) else str(cert)
                issuer = cert.get("issuer", "") if isinstance(cert, dict) else ""
                date = cert.get("issue_date", "") if isinstance(cert, dict) else ""
                p = doc.add_paragraph()
                r1 = p.add_run(cname)
                r1.bold = True
                if issuer or date:
                    p.add_run(f" — {issuer} ({date})")

    section_renderers = {
        "summary": docx_summary,
        "skills": docx_skills,
        "experiences": docx_experiences,
        "projects": docx_projects,
        "education": docx_education,
        "certifications": docx_certifications,
    }

    order = TEMPLATE_SECTION_ORDERS.get(template_name, TEMPLATE_SECTION_ORDERS["classic_ats"])
    for sec in order:
        renderer = section_renderers.get(sec)
        if renderer:
            renderer()

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
